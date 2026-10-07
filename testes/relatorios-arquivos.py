"""Verifica os arquivos reais de relatorios-e2e.mjs (openpyxl e PyMuPDF)."""
import json
import pathlib
import posixpath
import re
import warnings
import zipfile
import xml.etree.ElementTree as ET

import fitz
import openpyxl

PASTA = pathlib.Path(__file__).resolve().parents[1] / 'capturas' / 'relatorios'


def validar_pacote(caminho):
    with warnings.catch_warnings(record=True) as alertas:
        livro = openpyxl.load_workbook(caminho)
    assert not alertas, [str(a.message) for a in alertas]
    with zipfile.ZipFile(caminho) as pacote:
        nomes = set(pacote.namelist())
        for nome in nomes:
            if not nome.endswith(('.xml', '.rels')):
                continue
            root = ET.fromstring(pacote.read(nome))
            for el in root.iter():
                if 'rgb' in el.attrib:
                    assert re.fullmatch(r'[0-9A-Fa-f]{8}', el.attrib['rgb']), f'ARGB inválido para Excel: {nome}, {el.attrib}'
                if el.tag.endswith('}oneCellAnchor'):
                    assert 'editAs' not in el.attrib, 'editAs só é válido em twoCellAnchor'
            if nome.endswith('.rels'):
                base = posixpath.dirname(nome).replace('/_rels', '') if nome != '_rels/.rels' else ''
                for rel in root:
                    if rel.attrib.get('TargetMode') == 'External':
                        continue
                    destino = rel.attrib['Target']
                    destino = destino[1:] if destino.startswith('/') else posixpath.normpath(posixpath.join(base, destino))
                    assert destino in nomes, f'Relacionamento quebrado: {nome} -> {destino}'
        assert not any(n.startswith('xl/charts/') for n in nomes), 'sem gráficos XML montados manualmente'
    return livro


for nome in ('vazio', 'relatorio', 'filtrado', 'volume'):
    simples = validar_pacote(PASTA / f'{nome}-simples.xlsx')
    visual = validar_pacote(PASTA / f'{nome}-indicadores.xlsx')
    assert simples.sheetnames == ['Inventário atual', 'Movimentações', 'Leia-me']
    assert visual.sheetnames == ['Visão geral', 'Prioridades']
    for folha in simples:
        assert not folha._images and not folha._charts
        assert not folha.merged_cells, 'dados simples sem células mescladas'
    for folha in [simples['Inventário atual'], simples['Movimentações']]:
        assert folha.freeze_panes == 'A2', 'cabeçalho na primeira linha'
        assert folha['A1'].value in ['Técnico', 'Saída observada']
    for folha in visual:
        assert folha.max_row <= 30, 'resumo visual não cresce com milhares de linhas'
        assert len(folha._images) == 2
        assert folha.page_setup.fitToWidth == folha.page_setup.fitToHeight == 1
        assert folha['A1'].font.bold
        for imagem in folha._images:
            assert imagem.anchor.to.col > imagem.anchor._from.col
            assert imagem.anchor.to.row > imagem.anchor._from.row
    assert visual['Visão geral']['A10'].number_format == '0.0%'
    pdf = fitz.open(PASTA / f'{nome}.pdf')
    assert len(pdf) == 2, f'{nome}: PDF precisa ficar em duas páginas'
    texto = '\n'.join(p.get_text() for p in pdf)
    assert 'Visão geral' in texto and 'Prioridades da operação' in texto
    assert 'Exportação simples' in texto
    assert sum(len(p.get_image_info()) for p in pdf) == 4
    for pagina in pdf:
        assert f'{pagina.number + 1} / 2' in pagina.get_text()
        for b in pagina.get_text('blocks'):
            assert b[0] >= 0 and b[1] >= 0 and b[2] <= pagina.rect.width + 1 and b[3] <= pagina.rect.height + 1
        for info in pagina.get_image_info():
            assert info['bbox'][3] < pagina.rect.height - 50, 'gráficos não invadem notas ou rodapé'
    if nome == 'relatorio':
        assert [visual['Visão geral'][c].value for c in ['A6', 'E6', 'I6', 'A10', 'E10', 'I10']] == [33, 8, 12, 1 / 3, 2, 3]
        assert simples['Inventário atual']['E2'].value == '000123'
        assert simples['Inventário atual']['E2'].data_type == 's'
        assert any(c.data_type == 's' and str(c.value).startswith('=HYPERLINK') for row in simples['Inventário atual'] for c in row)
        assert simples['Inventário atual']['G2'].data_type == 'n'
        assert simples['Movimentações']['A2'].data_type == 'd'
        assert simples['Inventário atual'].max_row == 6
        assert simples['Movimentações'].max_row == 7
        assert 'Técnico sem peças 24' not in texto, 'resumo omite técnicos sem prioridade'
        assert 'BASE EXCLUIR' not in texto and 'IGNORADO EXCLUIR' not in texto
        assert 'Ana Exemplo' in texto
    elif nome == 'filtrado':
        assert visual['Visão geral']['A6'].value == 6
        assert simples['Inventário atual'].max_row == 2
        assert simples['Inventário atual']['A2'].value == 'Bruno Exemplo'
        assert 'Ana Exemplo' not in texto
    elif nome == 'volume':
        assert simples['Inventário atual'].max_row == 12001, '12.000 linhas preservadas no Excel simples'
        assert visual['Visão geral']['A6'].value == 606000
        assert visual['Visão geral']['E6'].value == 606000
        assert '8 de 120' in texto
        assert (PASTA / 'volume-indicadores.xlsx').stat().st_size < 1_000_000
    if nome in ('relatorio', 'volume'):
        for n, pagina in enumerate(pdf):
            pagina.get_pixmap(matrix=fitz.Matrix(1.6, 1.6)).save(PASTA / f'{nome}-pdf-{n + 1}.png')
    print(f'OK {nome}: simples completo, indicadores em 2 abas, PDF de 2 páginas, ARGB e relações válidos.')

# Nova importação precisa aparecer na nova exportação, sem atualizar a página manualmente.
importada = validar_pacote(PASTA / 'apos-importacao.xlsx')
esperado = json.loads((PASTA / 'apos-importacao.json').read_text())
assert importada['Visão geral']['A6'].value == esperado['estoque']
assert importada['Visão geral']['I6'].value == esperado['devolvidas']
print('OK nova importação: a nova exportação reflete o saldo e a devolução atualizados.')
